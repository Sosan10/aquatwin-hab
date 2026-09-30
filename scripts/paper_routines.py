import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Water Research / Environmental Modelling & Software — Manuscript Draft (Q1)")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Solorzano-Sanchez et al. (2026) — Gemelo Digital Limnológico para Alerta Temprana de HABs")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

def insert_table(doc, headers, data, caption, col_widths=None):
    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(14)
    cp.paragraph_format.space_after = Pt(4)
    cpr = cp.add_run(caption)
    cpr.font.name = "Times New Roman"
    cpr.font.size = Pt(9.5)
    cpr.font.bold = True
    cpr.font.color.rgb = RGBColor(15, 23, 42)
    
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    for idx, h in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    for r_idx, row_values in enumerate(data):
        row = table.rows[r_idx + 1]
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            r.font.name = "Times New Roman"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            if c_idx > 0 and (val.replace('.', '', 1).replace('-', '', 1).isdigit() or '%' in val or '±' in val or '<' in val):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(8)

def insert_figure(doc, fig_id, title, desc, ascii_diagram=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r_id = p.add_run(f"Figura {fig_id}. ")
    r_id.font.name = "Times New Roman"
    r_id.font.size = Pt(9.5)
    r_id.font.bold = True
    r_id.font.color.rgb = RGBColor(15, 23, 42)
    
    r_title = p.add_run(title)
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 41, 59)
    
    if ascii_diagram:
        box = doc.add_table(rows=1, cols=1)
        box.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = box.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
        bp = cell.paragraphs[0]
        bp.paragraph_format.space_before = Pt(0)
        bp.paragraph_format.space_after = Pt(0)
        bp.paragraph_format.line_spacing = 1.05
        br = bp.add_run(ascii_diagram)
        br.font.name = "Consolas"
        br.font.size = Pt(7.0)
        br.font.color.rgb = RGBColor(15, 23, 42)
        
    dp = doc.add_paragraph()
    dp.paragraph_format.space_before = Pt(4)
    dp.paragraph_format.space_after = Pt(12)
    dr = dp.add_run(desc)
    dr.font.name = "Times New Roman"
    dr.font.size = Pt(8.5)
    dr.font.italic = True
    dr.font.color.rgb = RGBColor(71, 85, 105)

def add_p(doc, text, bold_prefix=None, italic=False, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

print("Loaded base docx routines.")

"""
Script de generación del documento oficial GD.docx
Compila todo el código del Gemelo Digital (GD) de AquaTwin HAB:
- Gemelo Digital 3D (Three.js, simulación hidrodinámica, floraciones algales, partículas de Microcystis)
- Motor del Gemelo Digital (Frontend & Dashboard)
- Componentes de Visualización y Análisis del GD (CRISPDM, Fichas, Gráficas, Grafo)
- Algoritmos y Lógica Limnológica en TypeScript
- Motor y Backend Python del Gemelo Digital
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("AquaTwin HAB — Sistema de Gemelo Digital (GD) | Código Fuente")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139) # slate-500
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("AquaTwin HAB • Embalse Falling Creek (FCR) • Dataset fcr_oapat.csv • Confidencial")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(148, 163, 184) # slate-400

def create_cover_page(doc, file_catalog):
    total_lines = sum(item['lines'] for item in file_catalog)
    total_files = len(file_catalog)
    
    # Title Spacer
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(40)
    
    # Category tag
    p_cat = doc.add_paragraph()
    p_cat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cat = p_cat.add_run("SISTEMA DE MONITOREO Y PRONÓSTICO DE FLORACIONES ALGALES (HABs)")
    r_cat.font.name = "Segoe UI"
    r_cat.font.size = Pt(9.5)
    r_cat.font.bold = True
    r_cat.font.color.rgb = RGBColor(13, 148, 136) # teal-600
    
    # Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("AQUATWIN HAB\nGEMELO DIGITAL (GD)")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42) # slate-900
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Compilación Integral del Código Fuente del Gemelo Digital 3D,\nMotor Analítico Limnológico, Algoritmos Predictivos y Servicios de Telemetría")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(71, 85, 105) # slate-600
    
    # Metric summary box (table)
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["TOTAL ARCHIVOS", "LÍNEAS DE CÓDIGO", "DATASET ACTIVO", "STACK TECNOLÓGICO"]
    values = [f"{total_files} módulos", f"{total_lines:,} líneas", "fcr_oapat.csv (1,960 obs)", "Three.js / React / Python"]
    
    for i, col in enumerate(table.columns):
        col.width = Inches(1.8)
    
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_background(cell, "0F172A") # slate-900
        set_cell_margins(cell, top=120, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Segoe UI"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor(148, 163, 184)
        
    for i, v in enumerate(values):
        cell = table.cell(1, i)
        set_cell_background(cell, "F1F5F9") # slate-100
        set_cell_margins(cell, top=100, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(v)
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        
    # Spacer
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(30)
    
    # Summary list
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(12)
    r = p_desc.add_run("Alcance del Módulo de Gemelo Digital:")
    r.font.name = "Segoe UI"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    
    bullets = [
        ("Gemelo Digital 3D (DigitalTwin3DCanvas.tsx): ", "Simulación hidrodinámica Three.js con batimetría de alta resolución, modelo espectral de floraciones de fitoplancton (Chl-a, NDCI, Ficocianina), sistema de 2,400 partículas biológicas de Microcystis aeruginosa, oleaje físico multiharmónico, corrientes vectoriales y telemetría de boyas."),
        ("Motor del Gemelo Digital (GDMotorModule.tsx): ", "Dashboard de control limnológico en tiempo real, monitoreo de estratificación térmica, termoclina, alertas tempranas de toxicidad y modos de simulación."),
        ("Componentes Visuales CRISP-DM (components/gd/): ", "Ciclo completo de minería de datos, fichas técnicas de masa de agua, gráficas de dispersión, perfiles batimétricos y grafos de causalidad OAPAT."),
        ("Núcleo Algorítmico Limnológico (src/gd/): ", "Índices Carlson (TSI), gradientes térmicos Delta-T, detección de anomalías multivariadas, predicción temporal CNN-LSTM y reglas regulatorias expertas."),
        ("Motor Python y Servicios Backend (gd_python/): ", "Arquitectura de backend para asimilación de datos, explicabilidad XAI (SHAP), pipelines de entrenamiento y API de telemetría.")
    ]
    
    for bold_text, normal_text in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(2)
        rb = bp.add_run(bold_text)
        rb.font.name = "Segoe UI"
        rb.font.size = Pt(9.5)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(30, 41, 59)
        
        rn = bp.add_run(normal_text)
        rn.font.name = "Segoe UI"
        rn.font.size = Pt(9.5)
        rn.font.color.rgb = RGBColor(71, 85, 105)
        
    doc.add_page_break()

def create_table_of_contents(doc, file_catalog):
    hp = doc.add_paragraph()
    hrun = hp.add_run("Índice de Módulos del Gemelo Digital (GD)")
    hrun.font.name = "Segoe UI"
    hrun.font.size = Pt(18)
    hrun.font.bold = True
    hrun.font.color.rgb = RGBColor(15, 23, 42)
    hp.paragraph_format.space_after = Pt(14)
    
    table = doc.add_table(rows=1 + len(file_catalog), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    widths = [Inches(0.4), Inches(2.8), Inches(0.8), Inches(0.8), Inches(2.4)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
            
    # Headers
    h_titles = ["#", "Archivo / Módulo", "Líneas", "Tamaño", "Descripción"]
    for i, t in enumerate(h_titles):
        c = table.cell(0, i)
        set_cell_background(c, "0F172A")
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        p = c.paragraphs[0]
        r = p.add_run(t)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(241, 245, 249)
        if i in [0, 2, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
    # Rows
    for idx, item in enumerate(file_catalog):
        row_cells = table.rows[idx + 1].cells
        bg_color = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for c in row_cells:
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=60, bottom=60, left=70, right=70)
            
        # Col 0: Index
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(str(idx + 1))
        r0.font.name = "Segoe UI"
        r0.font.size = Pt(8)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(100, 116, 139)
        
        # Col 1: File name
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.add_run(item['path'])
        r1.font.name = "Consolas"
        r1.font.size = Pt(7.5)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(14, 116, 144) # cyan-700
        
        # Col 2: Lines
        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(f"{item['lines']:,}")
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor(51, 65, 85)
        
        # Col 3: Size
        p3 = row_cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(f"{item['size_kb']:.1f} KB")
        r3.font.name = "Segoe UI"
        r3.font.size = Pt(8)
        r3.font.color.rgb = RGBColor(51, 65, 85)
        
        # Col 4: Description
        p4 = row_cells[4].paragraphs[0]
        r4 = p4.add_run(item['desc'])
        r4.font.name = "Segoe UI"
        r4.font.size = Pt(7.5)
        r4.font.color.rgb = RGBColor(71, 85, 105)
        
    doc.add_page_break()

def add_file_code(doc, item, section_num):
    # Section Header Card
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(14)
    p_sec.paragraph_format.space_after = Pt(4)
    r_sec = p_sec.add_run(f"SECCIÓN {section_num}: {item['category'].upper()}")
    r_sec.font.name = "Segoe UI"
    r_sec.font.size = Pt(9.5)
    r_sec.font.bold = True
    r_sec.font.color.rgb = RGBColor(13, 148, 136) # teal-600
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run(item['path'])
    r_title.font.name = "Consolas"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    
    # Metadata info bar (table)
    m_table = doc.add_table(rows=1, cols=4)
    m_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.8), Inches(1.4), Inches(1.4), Inches(2.6)]
    for idx, width in enumerate(widths):
        m_table.rows[0].cells[idx].width = width
        
    meta_items = [
        ("Lenguaje / Tipo", item['lang']),
        ("Líneas", f"{item['lines']:,} líneas"),
        ("Tamaño", f"{item['size_kb']:.1f} KB"),
        ("Propósito", item['desc'])
    ]
    for idx, (label, val) in enumerate(meta_items):
        cell = m_table.cell(0, idx)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
        p = cell.paragraphs[0]
        rl = p.add_run(f"{label}: ")
        rl.font.name = "Segoe UI"
        rl.font.size = Pt(7.5)
        rl.font.bold = True
        rl.font.color.rgb = RGBColor(100, 116, 139)
        rv = p.add_run(val)
        rv.font.name = "Segoe UI"
        rv.font.size = Pt(7.5)
        rv.font.color.rgb = RGBColor(30, 41, 59)
        
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(4)
    
    # Read and insert code
    with open(item['abs_path'], 'r', encoding='utf-8', errors='replace') as fp:
        lines = fp.readlines()
        
    # We will format lines in paragraphs
    for line_idx, line in enumerate(lines, start=1):
        clean_line = line.rstrip('\r\n')
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        
        # Line number prefix
        r_num = p.add_run(f"{line_idx:4d} | ")
        r_num.font.name = "Consolas"
        r_num.font.size = Pt(7.0)
        r_num.font.color.rgb = RGBColor(148, 163, 184) # subtle grey
        
        # Code line
        r_code = p.add_run(clean_line if clean_line else " ")
        r_code.font.name = "Consolas"
        r_code.font.size = Pt(7.5)
        r_code.font.color.rgb = RGBColor(15, 23, 42) # slate-900
        
    doc.add_page_break()

def main():
    base_dir = os.path.abspath(r"c:\Users\crema\Downloads\aquatwin-hab---3d-digital-twin")
    output_docx = os.path.join(base_dir, "GD.docx")
    
    # Complete catalog of GD components and algorithms
    file_catalog = [
        # --- PARTE 1: GEMELO DIGITAL 3D ---
        {
            "category": "Gemelo Digital 3D (Simulación Visual Three.js)",
            "path": "src/components/DigitalTwin3DCanvas.tsx",
            "lang": "TypeScript (React / Three.js)",
            "desc": "Canvas 3D del Gemelo Digital con batimetría, floraciones algales, partículas de Microcystis y telemetría"
        },
        # --- PARTE 2: MOTOR DEL GEMELO DIGITAL (FRONTEND) ---
        {
            "category": "Motor del Gemelo Digital (Frontend)",
            "path": "src/components/GDMotorModule.tsx",
            "lang": "TypeScript (React / Tailwind)",
            "desc": "Panel de control limnológico, monitoreo de estratificación y diagnóstico determinista"
        },
        # --- PARTE 3: COMPONENTES VISUALES DEL GD ---
        {
            "category": "Visualización y Análisis CRISP-DM",
            "path": "src/components/gd/CRISPDM.tsx",
            "lang": "TypeScript (React)",
            "desc": "Fases completas de minería de datos y ciclo de vida de modelos del GD"
        },
        {
            "category": "Visualización y Análisis CRISP-DM",
            "path": "src/components/gd/Ficha.tsx",
            "lang": "TypeScript (React)",
            "desc": "Ficha técnica de la masa de agua, parámetros limnológicos y condiciones base"
        },
        {
            "category": "Visualización y Análisis CRISP-DM",
            "path": "src/components/gd/Graficas.tsx",
            "lang": "TypeScript (React)",
            "desc": "Series temporales limnológicas, perfiles de profundidad y correlaciones multivariadas"
        },
        {
            "category": "Visualización y Análisis CRISP-DM",
            "path": "src/components/gd/GraficasCRISPDM.tsx",
            "lang": "TypeScript (React)",
            "desc": "Gráficas de métricas de precisión, curvas ROC, matrices de confusión y pérdidas"
        },
        {
            "category": "Visualización y Análisis CRISP-DM",
            "path": "src/components/gd/GrafoOAPAT.tsx",
            "lang": "TypeScript (React)",
            "desc": "Visualización en grafo de causalidad y red de telemetría de boyas oceanográficas"
        },
        # --- PARTE 4: NÚCLEO MATEMÁTICO Y LIMNOLÓGICO DEL GD EN TYPESCRIPT ---
        {
            "category": "Núcleo Limnológico TypeScript (src/gd/)",
            "path": "src/gd/dominio.ts",
            "lang": "TypeScript",
            "desc": "Modelos de dominio, tipos de boyas, umbrales biológicos y estados del gemelo digital"
        },
        {
            "category": "Núcleo Limnológico TypeScript (src/gd/)",
            "path": "src/gd/indices.ts",
            "lang": "TypeScript",
            "desc": "Cálculo de índice Carlson TSI, gradiente térmico Delta-T, relación N:P y estabilidad hídrica"
        },
        {
            "category": "Núcleo Limnológico TypeScript (src/gd/)",
            "path": "src/gd/anomalias.ts",
            "lang": "TypeScript",
            "desc": "Detección explicable de anomalías multivariadas, picos atípicos y deriva sensorial"
        },
        {
            "category": "Núcleo Limnológico TypeScript (src/gd/)",
            "path": "src/gd/prediccion.ts",
            "lang": "TypeScript",
            "desc": "Motor de pronóstico a corto plazo (+24h, +48h, +72h) con explicabilidad exacta de factores"
        },
        {
            "category": "Núcleo Limnológico TypeScript (src/gd/)",
            "path": "src/gd/reglas.ts",
            "lang": "TypeScript",
            "desc": "Sistema experto de reglas limnológicas deterministas para clasificación de blooms"
        },
        {
            "category": "Núcleo Limnológico TypeScript (src/gd/)",
            "path": "src/gd/serie.ts",
            "lang": "TypeScript",
            "desc": "Generador de serie histórica continua, imputación y control de calidad de telemetría"
        },
        {
            "category": "Núcleo Limnológico TypeScript (src/gd/)",
            "path": "src/gd/crispdm.ts",
            "lang": "TypeScript",
            "desc": "Gestión estructurada de fases CRISP-DM y metadatos de modelos del Gemelo Digital"
        },
        # --- PARTE 5: BACKEND Y MOTOR PYTHON DEL GEMELO DIGITAL ---
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/app.py",
            "lang": "Python (FastAPI / Flask)",
            "desc": "API de servicios del Gemelo Digital, endpoints de telemetría, pronóstico y estado"
        },
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/aquatwin/dominio.py",
            "lang": "Python",
            "desc": "Modelo de dominio hidrológico y definiciones de parámetros biofísicos en Python"
        },
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/aquatwin/datos.py",
            "lang": "Python",
            "desc": "Ingesta, validación, transformación y alineación temporal del dataset real fcr_oapat.csv"
        },
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/aquatwin/explicabilidad.py",
            "lang": "Python",
            "desc": "Motor de explicabilidad XAI con atribución de factores causales de floraciones"
        },
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/aquatwin/motor/prediccion.py",
            "lang": "Python",
            "desc": "Modelos predictivos de machine learning (Ensemble, LSTM, Gradient Boosting)"
        },
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/aquatwin/motor/reglas.py",
            "lang": "Python",
            "desc": "Reglas limnológicas y regulatorias OMS/EPA para alertas de microcistina"
        },
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/aquatwin/motor/indices.py",
            "lang": "Python",
            "desc": "Cálculo de bio-índices y métricas de calidad de agua en Python"
        },
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/aquatwin/motor/anomalias.py",
            "lang": "Python",
            "desc": "Detección de anomalías en telemetría de boyas por Isolation Forest y Z-Score"
        },
        {
            "category": "Backend y Motor Python (gd_python/)",
            "path": "gd_python/aquatwin/vistas/gemelo.py",
            "lang": "Python",
            "desc": "Rutas y vistas de estado global del Gemelo Digital"
        }
    ]
    
    # Check and enrich catalog with size and line counts
    valid_catalog = []
    for item in file_catalog:
        abs_p = os.path.join(base_dir, item['path'].replace('/', os.sep))
        if os.path.exists(abs_p):
            item['abs_path'] = abs_p
            item['size_bytes'] = os.path.getsize(abs_p)
            item['size_kb'] = item['size_bytes'] / 1024.0
            with open(abs_p, 'r', encoding='utf-8', errors='replace') as fp:
                item['lines'] = len(fp.readlines())
            valid_catalog.append(item)
        else:
            print(f"Warning: File not found: {abs_p}")
            
    print(f"Total valid modules found: {len(valid_catalog)}")
    total_lines = sum(x['lines'] for x in valid_catalog)
    print(f"Total lines of code to compile: {total_lines:,}")
    
    doc = docx.Document()
    add_header_footer(doc)
    
    print("Generating cover page...")
    create_cover_page(doc, valid_catalog)
    
    print("Generating table of contents...")
    create_table_of_contents(doc, valid_catalog)
    
    print("Appending code files...")
    for idx, item in enumerate(valid_catalog, start=1):
        print(f"[{idx}/{len(valid_catalog)}] Processing {item['path']} ({item['lines']:,} lines)...")
        add_file_code(doc, item, idx)
        
    print(f"Saving final document to: {output_docx}...")
    doc.save(output_docx)
    print(f"Success! Document created successfully at: {output_docx}")
    file_size_mb = os.path.getsize(output_docx) / (1024 * 1024)
    print(f"File size: {file_size_mb:.2f} MB")

if __name__ == '__main__':
    main()
